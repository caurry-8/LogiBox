"""报告中心数据质量与覆盖率披露的回归测试（V3.5.0-A-03）。

验证报告能够回答：
- 数据总量 / 有效量 / 未参与量 / 覆盖率
- 没有可计算 CV 的有效 SKU 时，不能把「没有数据」写成「平均 CV 0.0000」
- 矩阵中被排除的 SKU 数量与原因
同时确认原有 7 章结构与示例数据的分析数值没有变化。
"""

from __future__ import annotations

import pandas as pd
import pytest
from docx import Document

from utils.abc_utils import ABCAnalyzer
from utils.data_store import DataStore
from utils.report_utils import generate_word_report
from utils.xyz_utils import XYZAnalyzer

REQUIRED_HEADINGS = [
    "一、数据概况",
    "二、ABC 分类分析",
    "三、XYZ 稳定性分析",
    "四、ABC × XYZ 交叉矩阵",
    "五、EOQ 计算结果",
    "六、安全库存与再订货点",
    "七、分析建议",
]


def _build_store(dataframe: pd.DataFrame, value_column: str, period_columns: list[str]):
    """用真实分析器填充 DataStore，不做任何 mock。"""
    store = DataStore()
    store.df = dataframe
    store.filename = "test_inventory.csv"

    abc = ABCAnalyzer(dataframe, value_column).analyze()
    store.set_analysis(
        "abc",
        {
            "dataframe": abc.dataframe,
            "counts": abc.counts,
            "contributions": abc.contributions,
            "value_column": value_column,
        },
    )

    xyz = XYZAnalyzer(dataframe, period_columns).analyze()
    store.set_analysis(
        "xyz",
        {
            "dataframe": xyz.dataframe,
            "counts": xyz.counts,
            "mean_cv": xyz.mean_cv,
            "status_counts": xyz.status_counts,
            "total_count": xyz.total_count,
            "period_columns": list(period_columns),
        },
    )
    return store, abc, xyz


def _rows(path) -> list[tuple[str, str]]:
    """收集报告所有「指标 / 结果」两列表格的行。"""
    document = Document(path)
    collected: list[tuple[str, str]] = []
    for table in document.tables:
        if not table.rows:
            continue
        if table.rows[0].cells[0].text != "指标":
            continue
        for row in table.rows[1:]:
            collected.append((row.cells[0].text, row.cells[1].text))
    return collected


def _value_of(path, label: str) -> str:
    matches = [value for key, value in _rows(path) if key == label]
    assert matches, f"报告中找不到指标「{label}」"
    return matches[0]


def _paragraphs(path) -> list[str]:
    return [p.text for p in Document(path).paragraphs]


def _headings(path) -> list[str]:
    return [
        p.text
        for p in Document(path).paragraphs
        if p.style.name.startswith("Heading")
    ]


# --------------------------------------------------------------------------- #
# 一、结构与前置条件
# --------------------------------------------------------------------------- #


def test_report_requires_data():
    store = DataStore()
    with pytest.raises(ValueError, match="没有可用于生成报告的数据"):
        generate_word_report("unused.docx", store)


def test_report_keeps_seven_sections(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _headings(path) == REQUIRED_HEADINGS


def test_report_without_xyz_analysis_says_not_completed(sample_dataframe, tmp_path):
    store = DataStore()
    store.df = sample_dataframe
    store.filename = "sample_inventory.csv"
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert "当前尚未完成 XYZ 稳定性分析。" in _paragraphs(path)


# --------------------------------------------------------------------------- #
# 二、数据概况披露
# --------------------------------------------------------------------------- #


def test_report_discloses_quality_status(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "数据质量状态") == "可分析"
    assert _value_of(path, "数据质量问题") == "阻断 0 / 警告 0 / 提示 0"


def test_report_discloses_sku_metrics(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "有效 SKU（SKU）") == "10"
    assert _value_of(path, "空值 SKU") == "0"
    assert _value_of(path, "重复 SKU") == "0"


def test_report_lists_quality_issues_for_dirty_data(tmp_path):
    dataframe = pd.DataFrame(
        {
            "SKU": ["A", "B"],
            "年消耗金额": [100.0, 50.0],
            "1月": [10, None],
            "2月": [20, None],
        }
    )
    store, _abc, _xyz = _build_store(dataframe, "年消耗金额", ["1月", "2月"])
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "数据质量状态") == "存在风险"
    bullets = [p for p in _paragraphs(path) if "存在缺失单元格" in p]
    assert bullets, "报告应列出缺失单元格问题"
    assert any("建议：" in bullet for bullet in bullets)


# --------------------------------------------------------------------------- #
# 三、XYZ 覆盖率与 mean_cv 口径
# --------------------------------------------------------------------------- #


def test_report_xyz_shows_coverage(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "总 SKU") == "10"
    assert _value_of(path, "有效分类 SKU") == "10"
    assert _value_of(path, "未参与分类 SKU") == "0"
    assert _value_of(path, "分类覆盖率") == "10 / 10（100.0%）"


def test_report_xyz_reports_partial_coverage(tmp_path):
    dataframe = pd.DataFrame(
        {
            "SKU": ["OK", "MISSING"],
            "年消耗金额": [100.0, 50.0],
            "1月": [10, None],
            "2月": [20, None],
        }
    )
    store, _abc, _xyz = _build_store(dataframe, "年消耗金额", ["1月", "2月"])
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "总 SKU") == "2"
    assert _value_of(path, "有效分类 SKU") == "1"
    assert _value_of(path, "未参与分类 SKU") == "1"
    assert _value_of(path, "分类覆盖率") == "1 / 2（50.0%）"


def test_report_xyz_mean_cv_uses_placeholder_when_no_valid_data(tmp_path):
    """全部 SKU 非法时必须写「暂无有效数据」，绝不能写成 0.0000。"""
    dataframe = pd.DataFrame(
        {
            "SKU": ["I1", "I2"],
            "年消耗金额": [100.0, 50.0],
            "1月": [float("inf"), float("inf")],
            "2月": [10.0, 20.0],
        }
    )
    store, _abc, xyz = _build_store(dataframe, "年消耗金额", ["1月", "2月"])
    assert xyz.mean_cv is None

    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "平均 CV") == "暂无有效数据"
    assert _value_of(path, "有效分类 SKU") == "0"


def test_report_xyz_mean_cv_keeps_numeric_value(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "平均 CV") == "0.1219"


def test_report_xyz_states_distribution_basis(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert "X / Y / Z 分布基于已完成有效分类的 SKU。" in _paragraphs(path)


# --------------------------------------------------------------------------- #
# 四、矩阵覆盖披露
# --------------------------------------------------------------------------- #


def test_report_matrix_discloses_exclusions(tmp_path):
    dataframe = pd.DataFrame(
        {
            "SKU": ["A", "B", "C", "D"],
            "年消耗金额": [400.0, 300.0, 200.0, 100.0],
            "1月": [10, None, float("inf"), 10],
            "2月": [10, None, 20, 10],
            "3月": [10, None, 30, 10],
        }
    )
    periods = ["1月", "2月", "3月"]
    store, abc, xyz = _build_store(dataframe, "年消耗金额", periods)

    from utils.matrix_utils import ABCXYZMatrixAnalyzer

    matrix = ABCXYZMatrixAnalyzer(
        abc.dataframe, xyz.dataframe, key_column="SKU", value_column="年消耗金额"
    ).analyze()
    store.set_analysis(
        "matrix",
        {
            "dataframe": matrix.dataframe,
            "cell_counts": matrix.cell_counts,
            "cell_values": matrix.cell_values,
            "cell_amounts": matrix.cell_amounts,
            "cell_sku_shares": matrix.cell_sku_shares,
            "key_column": matrix.key_column,
            "value_column": matrix.value_column,
            "total_value": matrix.total_value,
            "input_sku_count": matrix.input_sku_count,
            "excluded_sku_count": matrix.excluded_sku_count,
            "excluded_by_quality": matrix.excluded_by_quality,
        },
    )

    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "矩阵分析覆盖") == "2 / 4（50.0%）"
    assert _value_of(path, "未参与分析 SKU") == "2"
    assert "缺失数据 1" in _value_of(path, "未参与原因")
    assert "非法数据 1" in _value_of(path, "未参与原因")
    assert _value_of(path, "占比口径") == "基于已参与分析的 SKU"


# --------------------------------------------------------------------------- #
# 五、原有数值不回归
# --------------------------------------------------------------------------- #


def test_report_sample_values_unchanged(sample_dataframe, sample_period_columns, tmp_path):
    store, _abc, _xyz = _build_store(
        sample_dataframe, "年需求量", sample_period_columns
    )
    path = generate_word_report(str(tmp_path / "report.docx"), store)

    assert _value_of(path, "数据行数") == "10"
    assert _value_of(path, "字段数量") == "10"
    assert _value_of(path, "A 类 SKU") == "5"
    assert _value_of(path, "B 类 SKU") == "3"
    assert _value_of(path, "C 类 SKU") == "2"
    assert _value_of(path, "X 类 SKU") == "4"
    assert _value_of(path, "Y 类 SKU") == "6"
    assert _value_of(path, "Z 类 SKU") == "0"
